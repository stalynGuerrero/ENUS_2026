from pathlib import Path
import csv, math, json, shutil
from collections import defaultdict
from openpyxl import load_workbook
import xlsxwriter
root=Path.cwd(); out=root/'03_Encuesta/01_Tamaño_Muestra/Output'
source=out/'04_asignacion_sexo_edad10_departamento_regimen.xlsx'
w=load_workbook(source,read_only=True,data_only=True); s=w['2_Asignacion_Celda'];s.reset_dimensions(); rows=list(s.values)[1:]
assert len(rows)==1848
cells=defaultdict(list)
for r in rows: cells[r[1]].append(r)
assert len(cells)==33
assert sum(r[6] for r in rows)==50168746
assert all(r[6]>0 for r in rows)
backup=root/'revision_coherencia/respaldo/04_asignacion_sexo_edad10_departamento_regimen.xlsx'
if not backup.exists(): shutil.copy2(source,backup)
book=xlsxwriter.Workbook(str(source)); header=book.add_format({'bold':True,'bg_color':'#1F4E79','font_color':'white','font_name':'Arial'}); normal=book.add_format({'font_name':'Arial','num_format':'0'}); decimal=book.add_format({'font_name':'Arial','num_format':'0.000000'});
def sheet(name,headers):
 s=book.add_worksheet(name);s.write_row(0,0,headers,header);s.freeze_panes(1,0);s.set_column(0,len(headers)-1,20,normal);return s
p=sheet('0_Parametros',['parametro','valor'])
for i,item in enumerate([('P',.5),('Z',1.96),('ME',.0392),('DEFF',1.5),('Factor',1.4),('Minimo por celda',10),('Fuente poblacion','Matriz departamental anterior, conteos BDUA marzo 2026'),('Metodo','10 por celda + reparto proporcional del remanente + resto mayor; empate por orden de fila')],1):p.write_row(i,0,item)
d=sheet('1_Departamento_n',['region','departamento','poblacion','n_mas','n_deff','n_corregido','meta_sin_redondear','meta_final'])
a=sheet('2_Asignacion_Celda',['region','departamento','regimen','sexo','edad','zona','poblacion_celda','poblacion_depto','meta_depto','minimo','remanente_proporcional','entero','fraccion','orden_resto','meta_celda'])
check=sheet('3_Chequeo_Cuadre',['region','departamento','meta','asignado','diferencia'])
all_result=[]; summary=defaultdict(lambda:defaultdict(lambda:[0,0])); metas=[];cursor=2
for di,(dep,cr) in enumerate(cells.items(),2):
 N=sum(r[6] for r in cr); raw=937.5/(1+937.5/N)*1.4;target=math.ceil(raw);metas.append(target)
 assert len(cr)==56 and target>=560
 d.write_row(di-1,0,[cr[0][0],dep,N]);d.write_formula(di-1,3,"='0_Parametros'!B3^2*'0_Parametros'!B2*(1-'0_Parametros'!B2)/'0_Parametros'!B4^2",decimal,625)
 d.write_formula(di-1,4,f"=D{di}*'0_Parametros'!B5",decimal,937.5)
 d.write_formula(di-1,5,f'=E{di}/(1+E{di}/C{di})',decimal,937.5/(1+937.5/N))
 d.write_formula(di-1,6,f"=F{di}*'0_Parametros'!B6",decimal,raw);d.write_formula(di-1,7,f'=ROUNDUP(G{di},0)',normal,target)
 rem=target-560;quotas=[rem*r[6]/N for r in cr];base=[math.floor(q) for q in quotas];fractions=[q-b for q,b in zip(quotas,base)];left=rem-sum(base);order=sorted(range(56),key=lambda j:(-fractions[j],j));rank={j:k+1 for k,j in enumerate(order)};alloc=[10+base[j]+int(rank[j]<=left) for j in range(56)];assert sum(alloc)==target and min(alloc)>=10
 start=cursor;end=cursor+55
 for j,r in enumerate(cr):
  ri=cursor;ix=ri-1;a.write_row(ix,0,list(r[:6])+[r[6]])
  a.write_formula(ix,7,f"='1_Departamento_n'!C{di}",normal,N);a.write_formula(ix,8,f"='1_Departamento_n'!H{di}",normal,target);a.write_formula(ix,9,"='0_Parametros'!B7",normal,10)
  a.write_formula(ix,10,f'=(I{ri}-SUM($J${start}:$J${end}))*G{ri}/H{ri}',decimal,quotas[j]);a.write_formula(ix,11,f'=INT(K{ri})',normal,base[j]);a.write_formula(ix,12,f'=K{ri}-L{ri}',decimal,fractions[j])
  a.write_formula(ix,13,f'=COUNTIF($M${start}:$M${end},">"&M{ri})+COUNTIF($M${start}:M{ri},M{ri})',normal,rank[j]);a.write_formula(ix,14,f'=J{ri}+L{ri}+IF(N{ri}<=I{ri}-SUM($J${start}:$J${end})-SUM($L${start}:$L${end}),1,0)',normal,alloc[j]);cursor+=1
  record=dict(zip(['region','departamento','regimen','sexo','edad','zona','poblacion'],list(r[:6])+[r[6]]));record['meta']=alloc[j];all_result.append(record)
  for key in ['region','regimen','sexo','edad','zona']:
   summary[key][record[key]][0]+=r[6];summary[key][record[key]][1]+=alloc[j]
 check.write_row(di-1,0,[cr[0][0],dep]);check.write_formula(di-1,2,f"='1_Departamento_n'!H{di}",normal,target);check.write_formula(di-1,3,f"=SUM('2_Asignacion_Celda'!O{start}:O{end})",normal,sum(alloc));check.write_formula(di-1,4,f'=D{di}-C{di}',normal,0)
for k,name in [('sexo','4_Resumen_Sexo'),('edad','5_Resumen_Edad'),('zona','6_Resumen_Zona'),('regimen','7_Resumen_Regimen'),('region','8_Resumen_Region')]:
 sh=sheet(name,[k,'poblacion','meta','porcentaje']);col={'region':'A','regimen':'C','sexo':'D','edad':'E','zona':'F'}[k]
 for ix,(label,(pop,num)) in enumerate(summary[k].items(),1):
  sh.write(ix,0,label);sh.write_formula(ix,1,f'=SUMIF(\'2_Asignacion_Celda\'!{col}2:{col}1849,A{ix+1},\'2_Asignacion_Celda\'!G2:G1849)',normal,pop);sh.write_formula(ix,2,f'=SUMIF(\'2_Asignacion_Celda\'!{col}2:{col}1849,A{ix+1},\'2_Asignacion_Celda\'!O2:O1849)',normal,num);sh.write_formula(ix,3,f'=C{ix+1}/SUM(C2:C{len(summary[k])+1})',book.add_format({'font_name':'Arial','num_format':'0.0%'}),num/43173)
book.close();assert sum(metas)==43173
# Verify persisted formula caches and controls.
w=load_workbook(source,data_only=True);assert sum(r[14] for r in list(w['2_Asignacion_Celda'].values)[1:])==43173;assert all(r[4]==0 for r in list(w['3_Chequeo_Cuadre'].values)[1:]);assert not any(c.data_type=='e' for sh in w for row in sh for c in row)
(root/'revision_coherencia/asignacion_actualizada.json').write_text(json.dumps({'summary':dict(summary),'cells':all_result,'total':43173,'minimum':min(r['meta'] for r in all_result),'maximum':max(r['meta'] for r in all_result),'at_minimum':sum(r['meta']==10 for r in all_result)},ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps({'total':43173,'cells':len(all_result),'minimum':min(r['meta'] for r in all_result),'maximum':max(r['meta'] for r in all_result),'at_minimum':sum(r['meta']==10 for r in all_result),'summary':dict(summary)},ensure_ascii=False))
